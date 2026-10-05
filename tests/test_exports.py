from pathlib import Path
import json
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/annual-report-analysis/scripts'))
from calculate import evaluate
from export import display, export, prepare, word, workbook
from recovery import RecoveryInput, calculate_recovery
import scenarios
from test_quantitative import manufacturing_case, reverse_realized_case
from test_recovery import example
from workpaper import Workpaper


class ExportTests(unittest.TestCase):
    def test_long_reconciliation_basis_remains_complete_with_supported_row_heights(self):
        raw = json.loads((Path(__file__).resolve().parents[1]/'examples/catl-2025/workpaper.json').read_text(encoding='utf-8'))
        basis = '; '.join(f'{i}×original_source_cell_{i}; 千元整数取整假设，不代表审计重要性' for i in range(38)) + '\n\n保留末尾空格  '
        raw['reconciliations'][0]['basis'] = basis
        w = Workpaper.model_validate(raw)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'workbook.xlsx'
            workbook(w, prepare(w, evaluate(w)), path)
            with ZipFile(path) as archive:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                strings = [''.join(item.itertext()) for item in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
                sheets = ET.fromstring(archive.read('xl/workbook.xml')).find('x:sheets', ns)
                exported = {s.get('name'): ET.fromstring(archive.read(f'xl/worksheets/sheet{i}.xml'))
                            for i, s in enumerate(sheets, 1)}
                for sheet in exported.values():
                    for row in sheet.findall('x:sheetData/x:row', ns):
                        self.assertLessEqual(float(row.get('ht', '15')), 409)
                details = exported['TextDetails']
                parts = [strings[int(cell.find('x:v', ns).text)] for row in details.findall('x:sheetData/x:row', ns)[1:]
                         for cell in row.findall('x:c', ns) if cell.get('r').startswith('D')]
                self.assertEqual(''.join(parts), w.reconciliations[0].basis)
                main = exported['Reconciliations']
                link = main.find("x:hyperlinks/x:hyperlink[@ref='H2']", ns)
                self.assertEqual(link.get('location'), "'TextDetails'!D2")
                self.assertIn('TextDetails!D2:', strings[int(main.find(".//x:c[@r='H2']/x:v", ns).text)])
                self.assertIn('ABS(E2)<=F2', main.find(".//x:c[@r='G2']/x:f", ns).text)

    def test_cash_boundary_constraints_and_actual_bridges_are_visible(self):
        source = reverse_realized_case('unit_price')
        source['path_basis'] = 'counterfactual'
        source['periods'][0]['fixed_cash_cost'] = 500
        source['reverse']['target_cash'] = 50
        term = dict(source='Constructed threshold, not a debt agreement', available_at='2024-12-31',
                    definition='Constructed input condition', test_date='2025-12-31',
                    relation='at_least', threshold=0)
        source['contracts'] = [dict(term, label='Cash condition', metric='cash_end'),
                               dict(term, label='Profit condition', metric='ebitda'),
                               dict(term, label='Leverage condition', metric='debt_to_ebitda',
                                    relation='at_most', threshold=3),
                               dict(term, label='Unprojected condition', metric='cash_end', test_date='2025-06-30')]
        artifact = scenarios.run(source)
        self.assertEqual(artifact['reverse']['rows'][0]['ebitda'], -50)
        self.assertEqual(artifact['reverse']['contracts'][2]['status'], 'not_tested')
        w = Workpaper.model_validate({
            'mandate': dict(title='Boundary visibility', entity='Constructed', industry='manufacturing',
                            purpose='Read existing scenario results', period_start='2025-01-01',
                            period_end='2025-12-31', cutoff='2025-06-30', accounting_basis='Constructed',
                            scope='single entity', version='v2'),
            'sources': [], 'evidence': [], 'facts': [], 'findings': [], 'sections': [],
            'quantitative': [scenarios.to_workpaper_result(artifact, 'cash', 'Cash', [], [])],
        })
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_path = root/'input.json'
            input_path.write_text(w.model_dump_json(), encoding='utf-8')
            export(input_path, root/'output')
            self.assertEqual(json.loads((root/'output/quantitative-cash.json').read_text(encoding='utf-8')), artifact)
            report = Document(root/'output/report.docx')
            boundary = next(t for t in report.tables if t.cell(0, 0).text == '边界期间')
            self.assertEqual(boundary.cell(1, 1).text.replace('\u2011', '-'), '-50.0000')
            self.assertEqual(boundary.cell(1, 4).text, '未计算')
            conditions = [t for t in report.tables if t.cell(0, 0).text == '条件和日期']
            self.assertEqual(len(conditions), 2)
            self.assertIn('not_tested', conditions[1].cell(3, 2).text)
            self.assertIn('no projection at test date or ratio denominator is not meaningful', conditions[1].cell(3, 2).text)
            bridges = [t for t in report.tables if t.cell(0, 0).text == '期间与指标']
            self.assertEqual(len(bridges), 2)
            self.assertIn('remaining -450.0000', bridges[1].cell(1, 1).text.replace('\u2011', '-'))
            self.assertIn('not_supplied', bridges[1].cell(2, 2).text)
            with ZipFile(root/'output/workbook.xlsx') as archive:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                strings = [''.join(item.itertext()) for item in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
                sheets = ET.fromstring(archive.read('xl/workbook.xml')).find('x:sheets', ns)
                exported = {s.get('name'): ET.fromstring(archive.read(f'xl/worksheets/sheet{i}.xml'))
                            for i, s in enumerate(sheets, 1)}
                headings = [strings[int(c.find('x:v', ns).text)] for c in
                            exported['Q1AtBoundary'].findall('x:sheetData/x:row[@r="1"]/x:c', ns)]
                self.assertIn('ebit', headings)
                self.assertIn('ebitda', headings)
                self.assertIn('debt_to_ebitda', headings)
                conditions = exported['Q1BoundaryContracts']
                self.assertEqual(strings[int(conditions.find(".//x:c[@r='G4']/x:v", ns).text)], 'not_tested')
                self.assertIsNone(conditions.find(".//x:c[@r='F4']/x:v", ns))
                bridge = exported['Q1BoundaryActual']
                self.assertEqual(float(bridge.find(".//x:c[@r='G2']/x:v", ns).text), -450)
                self.assertIsNone(bridge.find(".//x:c[@r='G3']/x:v", ns))
                self.assertIn('Q1ActualBridge', exported)
            with ZipFile(root/'output/presentation.pptx') as archive:
                slides = [''.join(ET.fromstring(archive.read(name)).itertext()) for name in archive.namelist()
                          if name.startswith('ppt/slides/slide') and name.endswith('.xml')]
                visible = ''.join(''.join(slides).split())
                for text in ['EBITDA -50', '债务/EBITDA 未计算', 'Leverage condition', 'Unprojected condition',
                             'not_tested', 'outside_input_threshold', 'remaining -450', 'not_supplied',
                             'same_scope true', 'conflict true', '不重建剩余期间，也不重设现金路径']:
                    self.assertIn(''.join(text.split()), visible)

    def test_small_reconciliation_residuals_and_tolerance_remain_readable(self):
        raw = json.loads((Path(__file__).resolve().parents[1]/'examples/catl-2025/workpaper.json').read_text(encoding='utf-8'))
        context = {**raw['facts'][0]['context'], 'measure':'ratio', 'currency':None,
                   'physical_unit':None, 'scale':'1', 'aggregation':'ratio', 'start':None}
        for key, value in [('rounded_ratio', '1.568'), ('unrounded_ratio', '1.5680110721557214')]:
            raw['facts'].append(dict(id=key, label='Constructed rounding example', concept='test_ratio',
                                    value=value, context=context, evidence=[], state='assumption',
                                    note='Constructed display check, not issuer data'))
        for key, actual, expected in [('positive_residual', 'unrounded_ratio', 'rounded_ratio'),
                                      ('negative_residual', 'rounded_ratio', 'unrounded_ratio')]:
            raw['reconciliations'].append(dict(id=key, label='Constructed rounding comparison',
                actual=actual, expected=expected, tolerance='0.00005',
                basis='假设原披露到百分数两位，小数比率容差为0.00005；保留完整依据和真实非零残差。'))
        w = Workpaper.model_validate(raw)
        result = evaluate(w)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'workbook.xlsx'
            workbook(w, prepare(w, result), path)
            with ZipFile(path) as archive:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                sheet = ET.fromstring(archive.read('xl/worksheets/sheet4.xml'))
                styles = ET.fromstring(archive.read('xl/styles.xml')).find('x:cellXfs', ns)
                for row in [len(w.reconciliations), len(w.reconciliations)+1]:
                    cell = sheet.find(f".//x:c[@r='E{row}']", ns)
                    self.assertAlmostEqual(abs(float(cell.find('x:v', ns).text)), 0.0000110721557214)
                    self.assertEqual(styles[int(cell.get('s'))].get('numFmtId'), '0')
                    self.assertEqual(float(sheet.find(f".//x:c[@r='F{row}']/x:v", ns).text), 0.00005)
                    self.assertGreater(float(sheet.find(f".//x:row[@r='{row}']", ns).get('ht', '15')), 15)

    def test_percentage_point_change_is_distinct_from_relative_growth(self):
        context = dict(entity='Example insurer', scope='consolidated', aggregation='ratio',
                       basis='Constructed', measure='ratio')
        current = {**context, 'end':'2026-06-30'}
        raw = {
            'mandate': dict(title='Ratio changes', entity='Example insurer', industry='insurance',
                            purpose='Constructed comparison', period_start='2026-01-01', period_end='2026-06-30',
                            cutoff='2026-09-30', accounting_basis='Constructed', scope='consolidated', version='1'),
            'sources': [dict(id='source', title='Constructed ratios', url='example-source', published='2026-06-30')],
            'evidence': [dict(id='e', source='source', locator='Constructed table',
                             observation='Current 156.80%; previous 161.77%', reliability='Constructed example')],
            'facts': [dict(id='now', label='Current ratio', concept='capital_ratio', value='1.568',
                           context=current, evidence=['e']),
                      dict(id='before', label='Previous ratio', concept='capital_ratio', value='1.6177',
                           context={**context, 'end':'2025-12-31'}, evidence=['e'])],
            'calculations': [dict(id='delta', label='Absolute ratio change', op='difference',
                                  terms=[{'ref':'now'}, {'ref':'before'}],
                                  context={**current, 'physical_unit':'percentage_points', 'scale':'0.01'},
                                  definition='Current minus previous ratio', interpretation='Percentage points', period_rule='comparison'),
                             dict(id='growth', label='Relative change', op='growth',
                                  terms=[{'ref':'now'}, {'ref':'before'}], context={**current, 'scale':'0.01'},
                                  definition='Change divided by previous ratio', interpretation='Relative growth', period_rule='comparison')],
            'findings': [dict(id='change', title='Capital ratio change', question='How did the ratio change?',
                              conclusion='Absolute {{delta}}; relative {{growth}}', mechanism='Two distinct comparisons.',
                              status='supported', evidence=['delta', 'growth'], alternatives=[], changes_if='Update comparable inputs.')],
            'sections': [dict(title='Capital', findings=['change'], figures=['delta', 'growth'])],
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root/'input.json'
            source.write_text(json.dumps(raw, ensure_ascii=False), encoding='utf-8')
            export(source, root/'output')
            for filename, prefix in [('report.docx', 'word/document'), ('presentation.pptx', 'ppt/slides/slide')]:
                with ZipFile(root/'output'/filename) as archive:
                    text = ''.join(''.join(ET.fromstring(archive.read(name)).itertext())
                                   for name in archive.namelist() if name.startswith(prefix) and name.endswith('.xml'))
                if filename == 'report.docx':
                    self.assertIn('\u20114.97', text)
                    text = text.replace('\u2011', '-')
                self.assertIn('-4.97个百分点', ''.join(text.split()), filename)
                self.assertIn('-3.07%', ''.join(text.split()), filename)
            with ZipFile(root/'output'/'workbook.xlsx') as archive:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                sheet = ET.fromstring(archive.read('xl/worksheets/sheet3.xml'))
                self.assertAlmostEqual(float(sheet.find(".//x:c[@r='D2']/x:v", ns).text), -0.0497)
                delta = sheet.find(".//x:c[@r='E2']", ns)
                self.assertEqual(delta.find('x:f', ns).text, 'D2*100')
                self.assertAlmostEqual(float(delta.find('x:v', ns).text), -4.97)
                self.assertAlmostEqual(float(sheet.find(".//x:c[@r='E3']/x:v", ns).text), -0.0497/1.6177)
                styles = ET.fromstring(archive.read('xl/styles.xml'))
                fmt = styles.find('x:cellXfs', ns)[int(delta.get('s'))].get('numFmtId')
                code = next(f.get('formatCode') for f in styles.find('x:numFmts', ns) if f.get('numFmtId') == fmt)
                self.assertIn('个百分点', code)
                self.assertNotIn('%', code)

    def test_custom_snapshot_and_numeric_reference_reach_all_deliverables(self):
        context = dict(entity='Example issuer', scope='documented_bridge', end='2025-12-31',
                       aggregation='instant', basis='Constructed', measure='money', currency='CNY')
        raw = {
            'mandate': dict(title='Snapshot export', entity='Example issuer', industry='example',
                            purpose='Constructed snapshot display test', period_start='2025-01-01',
                            period_end='2025-12-31', cutoff='2026-10-03', accounting_basis='Constructed',
                            scope='single entity', version='1'),
            'sources': [dict(id='source', title='Constructed inputs', url='example-source', published='2025-12-31')],
            'evidence': [dict(id='bridge_evidence', source='source', locator='Example workpaper',
                             observation='Inputs 12 and 8', reliability='Constructed test only')],
            'facts': [],
            'findings': [dict(id='finding', title='Bridge', question='What is the snapshot result?',
                              conclusion='Bridge {{bridge_value}}', mechanism='Displayed source value.',
                              status='conditional', evidence=['bridge_value'], alternatives=[], changes_if='Rerun inputs.')],
            'sections': [dict(title='Bridge result', findings=['finding'], figures=['bridge_value'])],
            'quantitative': [dict(id='bridge', label='Custom bridge', method='custom_bridge_v1', as_of='2025-12-31',
                                  input_refs=[], evidence=['bridge_evidence'], assumptions=['Constructed inputs only'],
                                  limitations=['Snapshot display is not recalculation'],
                                  artifact={'input_snapshot': {'left': '12', 'right': '8'},
                                            'rows': [{'metric': 'Computed bridge', 'formula': '12 - 8', 'result': '4', 'unit': 'CNY',
                                                      'components': {'收入': '12', '支出': '8'}, 'conditions': ['核实', None]}]},
                                  figures=[dict(id='bridge_value', label='Computed bridge', row=0, field='result', context=context)])],
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for value, expected in [('4', '4.00 元人民币'), ('5', '5.00 元人民币'), (None, '未计算')]:
                with self.subTest(value=value):
                    raw['quantitative'][0]['artifact']['rows'][0]['result'] = value
                    input_path = root/'input.json'
                    input_path.write_text(json.dumps(raw, ensure_ascii=False), encoding='utf-8')
                    export(input_path, root/'output')
                    report = Document(root/'output'/'report.docx')
                    source_cell = report.tables[0].cell(1, 2).text
                    self.assertIn('bridge', source_cell)
                    self.assertNotIn('bridge_evidence', source_cell)
                    self.assertIn('Custom bridge [bridge]', [p.text for p in report.paragraphs])
                    for filename, prefix in [('report.docx', 'word/document'),
                                             ('workbook.xlsx', 'xl/sharedStrings'),
                                             ('presentation.pptx', 'ppt/slides/slide')]:
                        with ZipFile(root/'output'/filename) as archive:
                            text = ''.join(''.join(ET.fromstring(archive.read(name)).itertext())
                                           for name in archive.namelist() if name.startswith(prefix) and name.endswith('.xml'))
                        text = ''.join(text.split())
                        required_text = [expected, 'Computed bridge', 'bridge_evidence',
                                         'Constructed inputs only', 'Snapshot display is not recalculation']
                        if filename == 'workbook.xlsx':
                            required_text += ['12 - 8', 'CNY', '{"收入":"12","支出":"8"}', '["核实",null]']
                        else:
                            self.assertNotIn('components', text, filename)
                            summary = text.split('Custombridge', 1)[1]
                            self.assertIn('Computedbridge', summary, filename)
                            self.assertIn(''.join(expected.split()), summary, filename)
                        for required in required_text:
                            self.assertIn(''.join(required.split()), text, filename)
                        if value != '4':
                            self.assertNotIn('4.00元人民币', text, filename)
                    artifact = json.loads((root/'output'/'quantitative-bridge.json').read_text(encoding='utf-8'))
                    self.assertEqual(artifact, raw['quantitative'][0]['artifact'])
        raw['findings'][0]['conclusion'] = '{{unknown_result}}'
        w = Workpaper.model_validate(raw)
        with self.assertRaisesRegex(ValueError, 'unknown numeric token'):
            prepare(w, evaluate(w))
        raw['quantitative'][0]['figures'][0]['field'] = 'unknown_field'
        with self.assertRaisesRegex(ValueError, 'unknown artifact row or field'):
            Workpaper.model_validate(raw)

    def test_methods_and_evidence_status_remain_visible_in_all_deliverables(self):
        method = 'Example corporate methodology 2026-10-03'
        for status in ('conditional', 'unresolved'):
            w = Workpaper.model_validate({
                'mandate': dict(title='Issuer credit review', entity='Example issuer', industry='manufacturing',
                                purpose='Issuer credit recommendation', period_start='2025-01-01',
                                period_end='2025-12-31', cutoff='2026-10-03', accounting_basis='CAS',
                                scope='consolidated', version='1', methods=[method]),
                'sources': [], 'evidence': [], 'facts': [],
                'findings': [dict(id='rating', title='Credit recommendation', question='What supports the recommendation?',
                                  conclusion='The recommendation depends on financing evidence.', mechanism='Funding availability.',
                                  status=status, evidence=[], alternatives=[], changes_if='Reassess if financing is withdrawn.')],
                'sections': [dict(title='Credit recommendation', findings=['rating'])],
            })
            with tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                data = prepare(w, evaluate(w))
                word(data, root/'report.docx')
                workbook(w, data, root/'workbook.xlsx')
                (root/'presentation-data.json').write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
                script = Path(__file__).resolve().parents[1]/'skills/annual-report-analysis/scripts/export_pptx.cjs'
                subprocess.run(['node', str(script), str(root/'presentation-data.json'), str(root/'presentation.pptx')], check=True)
                for filename, members in [('report.docx', ['word/document.xml']),
                                          ('workbook.xlsx', ['xl/sharedStrings.xml']),
                                          ('presentation.pptx', None)]:
                    with self.subTest(status=status, file=filename), ZipFile(root/filename) as archive:
                        visible = members or [name for name in archive.namelist()
                                              if name.startswith('ppt/slides/slide') and name.endswith('.xml')]
                        text = '\n'.join(''.join(ET.fromstring(archive.read(name)).itertext()) for name in visible)
                        self.assertIn(method, text)
                        self.assertIn(status, text)

    def test_small_cny_and_direct_finding_numbers_keep_their_meaning(self):
        raw = json.loads((Path(__file__).resolve().parents[1]/'examples/catl-2025/workpaper.json').read_text(encoding='utf-8'))
        raw['findings'][0]['conclusion'] = 'Cash bridge {{cfo_bridge}}'
        raw['findings'][0]['evidence'] = ['cfo_bridge']
        w = Workpaper.model_validate(raw)
        ctx = w.facts[0].context.model_copy(update={'currency':'CNY', 'measure':'money'})
        for scale, expected in [(1000000, '2,500.00 万元人民币'), (1, '25.00 元人民币')]:
            with self.subTest(scale=scale):
                self.assertEqual(display(25, ctx.model_copy(update={'scale':scale})), expected)
        self.assertEqual(display(-25, ctx.model_copy(update={'currency':'USD', 'scale':1000000})), '-25.00 百万USD')
        self.assertEqual(prepare(w, evaluate(w))['findings'][0]['figure_refs'], ['cfo_bridge'])
        self.assertEqual(w.findings[0].conclusion, 'Cash bridge {{cfo_bridge}}')

    def test_cash_filter_and_print_view_include_both_periods(self):
        source = manufacturing_case()
        source['periods'].append({**source['periods'][0], 'start':'2026-01-01', 'end':'2026-12-31'})
        artifact = scenarios.run(source)
        w = Workpaper.model_validate({
            'mandate': dict(title='Two periods', entity='Constructed', industry='manufacturing', purpose='Print cash path',
                            period_start='2025-01-01', period_end='2026-12-31', cutoff='2024-12-31',
                            accounting_basis='Constructed', scope='single entity', version='v2'),
            'sources': [], 'evidence': [], 'facts': [], 'findings': [], 'sections': [],
            'quantitative': [scenarios.to_workpaper_result(artifact, 'cash', 'Cash', [], [])],
        })
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'workbook.xlsx'
            workbook(w, prepare(w, evaluate(w)), path)
            with ZipFile(path) as z:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                book = ET.fromstring(z.read('xl/workbook.xml'))
                index = next(i for i, s in enumerate(book.find('x:sheets', ns), 1) if s.get('name') == 'Q1Cash')
                xml = ET.fromstring(z.read(f'xl/worksheets/sheet{index}.xml'))
                self.assertEqual(xml.find('x:autoFilter', ns).get('ref'), 'A1:AG3')
                print_formula = [x.text for x in xml.findall('.//x:f', ns) if x.text in ('V2', 'V3')]
                self.assertEqual(print_formula, ['V2', 'V3'])
                for row, result in enumerate(artifact['rows'], 2):
                    self.assertEqual(float(xml.find(f".//x:c[@r='V{row}']/x:v", ns).text), result['cash_end'])
                    for col, expected in [('AD', 200), ('AE', 90 if row == 2 else 94), ('AF', 0), ('AG', 50)]:
                        self.assertEqual(float(xml.find(f".//x:c[@r='{col}{row}']/x:v", ns).text), expected)
                area = next(x.text for x in book.findall('x:definedNames/x:definedName', ns)
                            if x.get('name') == '_xlnm.Print_Area' and x.text.startswith('Q1Cash!'))
                self.assertEqual(area, 'Q1Cash!$A$7:$C$39')

    def test_explicit_times_and_default_percentage_are_distinct_formats(self):
        raw = json.loads((Path(__file__).resolve().parents[1]/'examples/catl-2025/workpaper.json').read_text(encoding='utf-8'))
        calc = next(c for c in raw['calculations'] if c['op'] == 'ratio')
        calc['context']['physical_unit'] = 'times'
        w = Workpaper.model_validate(raw)
        c = next(c for c in w.calculations if c.id == calc['id'])
        self.assertEqual(display('43.5978', c.context), '43.60 倍')
        self.assertEqual(display('0.435978', c.context.model_copy(update={'physical_unit':None})), '43.60%')
        result = evaluate(w)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'workbook.xlsx'
            workbook(w, prepare(w, result), path)
            with ZipFile(path) as z:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                xml = ET.fromstring(z.read('xl/worksheets/sheet3.xml'))
                row = next(i for i, record in enumerate(w.calculations, 2) if record.id == c.id)
                cell = xml.find(f".//x:c[@r='E{row}']", ns)
                style = ET.fromstring(z.read('xl/styles.xml'))
                numfmt = style.find('x:cellXfs', ns)[int(cell.get('s'))].get('numFmtId')
                code = next(f.get('formatCode') for f in style.find('x:numFmts', ns) if f.get('numFmtId') == numfmt)
                self.assertIn('" 倍"', code)
                self.assertNotIn('%', code)
                self.assertEqual(float(cell.find('x:v', ns).text), float(result['calculations'][row-2]['value']))

    def test_historical_cash_artifact_cannot_be_relabelled_as_v2(self):
        raw = json.loads((Path(__file__).resolve().parents[1]/'tests/fixtures/legacy-scenario-workpaper.json').read_text(encoding='utf-8'))
        raw['quantitative'] = raw['quantitative'][:1]
        w = Workpaper.model_validate(raw)
        with self.assertRaisesRegex(ValueError, 'historical model'):
            prepare(w, evaluate(w))
        raw['quantitative'][0]['method'] = 'operating_cash_scenario_v2'
        raw['quantitative'][0]['artifact']['method'] = 'operating_cash_scenario_v2'
        w = Workpaper.model_validate(raw)
        with self.assertRaisesRegex(ValueError, 'unit_cost_of_sales'):
            prepare(w, evaluate(w))

    def test_manufacturing_cost_decomposition_reaches_cash_formulas(self):
        for inventory_change, ending_inventory, cash in [(0, 60, 100), (10, 90, 110), (-10, 70, 110)]:
            with self.subTest(inventory_depreciation_change=inventory_change):
                source = manufacturing_case()
                if inventory_change:
                    source['periods'][0].update(depreciation=30, depreciation_in_cost_of_sales=30,
                        inventory_cash_conversion=20, inventory_depreciation_change=inventory_change,
                        dio=ending_inventory*365/120, source='Constructed manufacturing D&A and payroll cash bridge')
                artifact = scenarios.run(source)
                w = Workpaper.model_validate({
                    'mandate': dict(title='Manufacturing cash', entity='Constructed', industry='manufacturing',
                                    purpose='Independent direct cash bridge', period_start='2025-01-01',
                                    period_end='2025-12-31', cutoff='2024-12-31', accounting_basis='Constructed',
                                    scope='single entity', version='v2'),
                    'sources': [], 'evidence': [], 'facts': [], 'findings': [], 'sections': [],
                    'quantitative': [scenarios.to_workpaper_result(artifact, 'manufacturing', 'Manufacturing', [], [])],
                })
                with tempfile.TemporaryDirectory() as directory:
                    root = Path(directory)
                    data = prepare(w, evaluate(w))
                    workbook(w, data, root/'workbook.xlsx')
                    word(data, root/'report.docx')
                    with ZipFile(root/'workbook.xlsx') as z:
                        ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        sheets = ET.fromstring(z.read('xl/workbook.xml')).find('x:sheets', ns)
                        index = next(i for i, s in enumerate(sheets, 1) if s.get('name') == 'Q1Cash')
                        xml = ET.fromstring(z.read(f'xl/worksheets/sheet{index}.xml'))
                    for ref, expected in [('L2', 100), ('M2', 40), ('V2', cash)]:
                        self.assertAlmostEqual(float(xml.find(f".//x:c[@r='{ref}']/x:v", ns).text), expected)
                    for ref, formula in [('D2', "B2-C2+'Q1Drivers'!R2-'Q1Drivers'!G2"),
                                         ('L2', "C2+K2-'Q1Inputs'!B8-'Q1Drivers'!S2-'Q1Drivers'!R2-'Q1Drivers'!T2"),
                                         ('P2', "I2+E2+'Q1Drivers'!T2-O2")]:
                        self.assertIn(formula, xml.find(f".//x:c[@r='{ref}']/x:f", ns).text)
                    self.assertIn("COUNT('Q1Drivers'!A2:T2)=20", xml.find(".//x:c[@r='V2']/x:f", ns).text)
                    with ZipFile(root/'report.docx') as z:
                        text = ' '.join(ET.fromstring(z.read('word/document.xml')).itertext())
                    self.assertIn(f'{cash:,.4f}', text)
                    self.assertIn('Supplier purchases exclude internal production cash conversion', text)

    def test_incomplete_ratio_or_growth_remains_unavailable_in_workbook(self):
        for operation in ('ratio', 'growth'):
            with self.subTest(operation=operation):
                raw = json.loads((Path(__file__).resolve().parents[1]/'examples/catl-2025/workpaper.json').read_text(encoding='utf-8'))
                calculation = next(c for c in raw['calculations'] if c['op'] == operation)
                calculation['terms'] = calculation['terms'][:1]
                w = Workpaper.model_validate(raw)
                result = evaluate(w)
                row = next(i for i, c in enumerate(w.calculations, 2) if c.id == calculation['id'])
                computed = result['calculations'][row-2]
                self.assertEqual(computed['status'], 'not_calculated')
                self.assertEqual(computed['reason'], 'operation requires two inputs')
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory)/'workbook.xlsx'
                    workbook(w, prepare(w, result), path)
                    with ZipFile(path) as z:
                        ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        xml = ET.fromstring(z.read('xl/worksheets/sheet3.xml'))
                    cell = xml.find(f".//x:c[@r='D{row}']", ns)
                    self.assertEqual(cell.find('x:f', ns).text, 'NA()')
                    self.assertEqual(cell.find('x:v', ns).text, '#N/A')

    def test_monetary_reconciliation_keeps_a_real_cent_difference(self):
        ctx = dict(entity='A', scope='consolidated', start='2025-01-01', end='2025-12-31',
                   aggregation='flow', basis='test', measure='money', currency='CNY')
        for right, residual, status in [('200000000000.22', '0.00', 'within_input_tolerance'),
                                        ('200000000000.23', '0.01', 'unexplained_difference')]:
            with self.subTest(right=right):
                w = Workpaper.model_validate({
                    'mandate': dict(title='Decimal precision', entity='A', industry='synthetic', purpose='Precision test',
                                    period_start='2025-01-01', period_end='2025-12-31', cutoff='2026-10-02',
                                    accounting_basis='test', scope='consolidated', version='test'),
                    'sources': [], 'evidence': [],
                    'facts': [dict(id=k, label=k, concept=k, value=v, context=ctx, evidence=[],
                                   state='assumption', note='Synthetic precision test')
                              for k, v in [('left', '100000000000.11'), ('right', right), ('total', '300000000000.33')]],
                    'calculations': [dict(id='sum', label='Sum', op='sum', terms=[{'ref':'left'}, {'ref':'right'}],
                                          context=ctx, definition='left+right', interpretation='Synthetic')],
                    'reconciliations': [dict(id='check', label='Exact', actual='sum', expected='total', tolerance='0',
                                             basis='Exact decimal inputs, two decimal places')],
                    'findings': [], 'sections': [],
                })
                result = evaluate(w)
                self.assertEqual(result['reconciliations'][0]['residual'], residual)
                self.assertEqual(result['reconciliations'][0]['status'], status)
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory)/'workbook.xlsx'
                    workbook(w, prepare(w, result), path)
                    with ZipFile(path) as z:
                        ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        sheets = ET.fromstring(z.read('xl/workbook.xml')).find('x:sheets', ns)
                        index = next(i for i, s in enumerate(sheets, 1) if s.get('name') == 'Reconciliations')
                        xml = ET.fromstring(z.read(f'xl/worksheets/sheet{index}.xml'))
                    cell = xml.find(".//x:c[@r='E2']", ns)
                    self.assertIn('ROUND(', cell.find('x:f', ns).text)
                    self.assertEqual(float(cell.find('x:v', ns).text), float(residual))
                    self.assertEqual(xml.find(".//x:c[@r='J2']/x:v", ns).text, '2')

    def test_procedure_and_request_numbers_are_resolved_in_both_documents(self):
        raw = json.loads((Path(__file__).resolve().parents[1]/'examples/catl-2025/workpaper.json').read_text(encoding='utf-8'))
        raw['procedures'][0]['result'] = 'Cash bridge {{cfo_bridge}}'
        raw['requests'][0]['request'] = 'Explain {{cfo_bridge}}'
        w = Workpaper.model_validate(raw)
        data = prepare(w, evaluate(w))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            word(data, root/'report.docx')
            workbook(w, data, root/'workbook.xlsx')
            for file, part in [('report.docx', 'word/document.xml'), ('workbook.xlsx', 'xl/sharedStrings.xml')]:
                with ZipFile(root/file) as z:
                    text = ' '.join(ET.fromstring(z.read(part)).itertext())
                self.assertIn('Cash bridge 1,332.20', text)
                self.assertIn('Explain 1,332.20', text)
                self.assertNotIn('{{cfo_bridge}}', text)
        raw['procedures'][0]['result'] = 'Unknown {{typo_cfo_missing}}'
        w = Workpaper.model_validate(raw)
        with self.assertRaisesRegex(ValueError, 'unknown numeric token'):
            prepare(w, evaluate(w))

    def test_recovery_results_reach_report_and_editable_workbook(self):
        result = calculate_recovery(RecoveryInput.model_validate(example()))
        w = Workpaper.model_validate({
            'mandate': dict(title='Constructed recovery', entity='Borrower A', industry='synthetic',
                            purpose='Verify recoveries survive export', period_start='2026-01-01',
                            period_end='2026-10-02', cutoff='2026-10-02', accounting_basis='Scenario',
                            scope='single entity', version='test'),
            'sources': [dict(id='s', title='Synthetic fixture', url='test_recovery.py', published='2026-10-02')],
            'evidence': [dict(id=k, source='s', locator='example()', observation='Constructed inputs',
                              reliability='Mathematical example only') for k in result.evidence],
            'facts': [], 'findings': [], 'sections': [], 'quantitative': [result],
        })
        data = prepare(w, evaluate(w))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            word(data, root/'report.docx')
            workbook(w, data, root/'workbook.xlsx')
            with ZipFile(root/'report.docx') as z:
                text = ' '.join(ET.fromstring(z.read('word/document.xml')).itertext())
            self.assertIn('Secured loan', text)
            self.assertIn('85.00', text)
            self.assertIn('85.00%', text)
            self.assertNotIn('已执行程序与待核工作', text)
            with ZipFile(root/'workbook.xlsx') as z:
                ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                sheets = ET.fromstring(z.read('xl/workbook.xml')).find('x:sheets', ns)
                recovery_index = next(i for i, s in enumerate(sheets, 1) if s.get('name') == 'Q1Recovery')
                xml = ET.fromstring(z.read(f'xl/worksheets/sheet{recovery_index}.xml'))
                cell = xml.find(".//x:c[@r='E2']", ns)
                self.assertEqual(cell.find('x:f', ns).text, 'IF(COUNT(C2:D2)=2,C2+D2,NA())')
                self.assertEqual(float(cell.find('x:v', ns).text), 85)


if __name__ == '__main__':
    unittest.main()

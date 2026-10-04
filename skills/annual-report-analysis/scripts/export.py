"""Calculate a workpaper once and export its editable research deliverables."""
import argparse
import hashlib
import json
import re
import subprocess
import platform
from datetime import datetime
from decimal import Decimal
from importlib.metadata import version
from pathlib import Path

import xlsxwriter
from xlsxwriter.utility import xl_col_to_name
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from calculate import context_error, evaluate, excel_expression, reconciliation_matches
from scenarios import Scenario
from workpaper import Workpaper


def unit(ctx):
    if ctx.measure == "money" and ctx.scale == 1000000:
        return f"百万{ctx.currency}"
    if ctx.measure == "ratio":
        return "倍" if ctx.physical_unit == "times" else "个百分点" if ctx.physical_unit == "percentage_points" else "%"
    else:
        measure = ctx.currency if ctx.measure == "money" else ctx.physical_unit if ctx.measure == "count" else ctx.measure
    return f"{measure} × {ctx.scale}" if ctx.scale != 1 else measure


def display(value, ctx):
    if value is None:
        return "未计算"
    value = Decimal(str(value))
    if ctx.measure == "ratio" and ctx.physical_unit == "times":
        return f"{value*ctx.scale:,.2f} 倍"
    if ctx.measure == "ratio" and ctx.physical_unit == "percentage_points":
        return f"{value*ctx.scale*100:,.2f} 个百分点"
    if ctx.measure == "money" and ctx.currency == "CNY":
        amount = value * ctx.scale
        divisor, label = (Decimal("1e8"), "亿元人民币") if abs(amount) >= Decimal("1e8") else (Decimal("1e4"), "万元人民币") if abs(amount) >= Decimal("1e4") else (Decimal(1), "元人民币")
        return f"{amount / divisor:,.2f} {label}"
    return f"{value*ctx.scale:.2%}" if ctx.measure == "ratio" else f"{value:,.2f} {unit(ctx)}"


def prepare(w, result):
    if any(q.method == "operating_cash_scenario_v1" for q in w.quantitative):
        raise ValueError("operating_cash_scenario_v1 is a historical model; replay/export it with its frozen code version or supply documented v2 cost decomposition")
    for q in w.quantitative:
        if q.method == "operating_cash_scenario_v2":
            if q.artifact.get("method") != q.method:
                raise ValueError("cash artifact method must match its workpaper method")
            Scenario.model_validate(q.artifact["input_snapshot"])
    figures = {}
    for f in w.facts:
        figures[f.id] = {"id": f.id, "label": f.label, "value": str(f.value) if f.value is not None else None,
                         "display": display(f.value, f.context), "context": f.context.model_dump(mode="json"),
                         "evidence": f.evidence, "note": f.note}
    computed = {r["id"]: r for r in result["calculations"]}
    for c in w.calculations:
        r = computed[c.id]
        evidence = sorted({e for t in c.terms for e in figures[t.ref]["evidence"]})
        figures[c.id] = {"id": c.id, "label": c.label, "value": r["value"],
                         "display": display(r["value"], c.context), "context": c.context.model_dump(mode="json"),
                         "evidence": evidence, "note": r["reason"] or c.interpretation}
    for q in w.quantitative:
        evidence = sorted(set(q.evidence) | {e for ref in q.input_refs for e in figures[ref]["evidence"]})
        for f in q.figures:
            value = q.artifact["rows"][f.row][f.field]
            figures[f.id] = {"id": f.id, "label": f.label, "value": str(value) if value is not None else None,
                             "display": display(value, f.context), "context": f.context.model_dump(mode="json"),
                             "evidence": evidence, "note": f"{q.id} artifact.rows[{f.row}].{f.field}；已运行结果快照，导出未重新计算或校验"}

    def resolve(text):
        def replace(match):
            key = match.group(1)
            if key not in figures:
                raise ValueError(f"unknown numeric token {{{{{key}}}}}")
            return figures[key]["display"]
        return re.sub(r"\{\{([^{}]+)\}\}", replace, text)

    data = w.model_dump(mode="json")
    unknown_dates = [s.id for s in w.sources if s.published is None]
    if unknown_dates:
        data["mandate"]["limitations"].append("来源 "+", ".join(unknown_dates)+" 的正式公布日期未核验；不得据此证明在信息截止日前已公开。当前披露内容可以继续分析，历史时点结论须补充公告时间证据。")
    narrative_fields = {
        "findings": ("title", "question", "conclusion", "mechanism", "changes_if", "alternatives"),
        "procedures": ("purpose", "assertions", "population", "selection", "steps", "result"),
        "requests": ("request", "reason", "close_when"),
        "sections": ("title",),
        "presentation": ("title", "body"),
        "quantitative": ("label", "assumptions", "limitations"),
    }
    for group, keys in narrative_fields.items():
        for record in data[group]:
            if group in ("findings", "presentation"):
                tokens = re.findall(r"\{\{([^{}]+)\}\}", "\n".join(x for key in keys for x in (record[key] if isinstance(record[key], list) else [record[key]])))
                record["figure_refs"] = list(dict.fromkeys([*tokens, *(ref for ref in record.get("evidence", []) if ref in figures)]))
            for key in keys:
                value = record[key]
                record[key] = [resolve(x) for x in value] if isinstance(value, list) else resolve(value)
    for key in ("title", "purpose"):
        data["mandate"][key] = resolve(data["mandate"][key])
    data["mandate"]["limitations"] = [resolve(x) for x in data["mandate"]["limitations"]]
    data["figures"] = figures
    data["results"] = result
    return data


def workbook(w, data, path):
    with xlsxwriter.Workbook(path, {"strings_to_formulas": False, "strings_to_urls": False}) as book:
        book.set_properties({"title": data["mandate"]["title"], "comments": f"Workpaper {w.mandate.version}"})
        head = book.add_format({"bold": True, "bg_color": "#213B45", "font_color": "white", "text_wrap": True})
        wrap = book.add_format({"text_wrap": True, "valign": "top"})
        number = book.add_format({"num_format": "#,##0.00;[Red](#,##0.00)"})
        percent = book.add_format({"num_format": "0.00%;[Red](0.00%)"})
        times = book.add_format({"num_format": '#,##0.00" 倍";[Red](#,##0.00" 倍")'})
        points = book.add_format({"num_format": '0.00" 个百分点";[Red](0.00" 个百分点")'})

        def sheet(name, headers, rows, widths=None):
            ws = book.add_worksheet(name)
            ws.write_row(0, 0, headers, head)
            ws.freeze_panes(1, 2)
            ws.set_row(0, 32)
            for rowno, row in enumerate(rows, 1):
                for col, item in enumerate(row):
                    ws.write(rowno, col, json.dumps(item, ensure_ascii=False) if isinstance(item, (dict, list)) else item, wrap)
                if name in ("Readme", "Reconciliations"):
                    lines = max(sum(1 + sum(2 if ord(char) > 255 else 1 for char in line) // (widths[col]-2)
                                    for line in str(item).split("\n")) for col, item in enumerate(row))
                    ws.set_row(rowno, 15 * lines + 3)
            ws.autofilter(0, 0, max(1, len(rows)), len(headers)-1)
            for col in range(len(headers)):
                ws.set_column(col, col, widths[col] if widths else 28)
            if name in ("Readme", "Calculations", "Reconciliations"):
                ws.set_landscape()
                ws.set_paper(9 if name == "Readme" else 8)
                ws.fit_to_pages(1, 0)
                ws.repeat_rows(0)
                if name != "Readme":
                    ws.set_margins(left=0.25, right=0.25)
            return ws

        sheet("Readme", ["字段 / field", "内容 / value"], [
            ["任务", data["mandate"]["title"]], ["版本", w.mandate.version], ["信息截止", str(w.mandate.cutoff)],
            ["范围", w.mandate.scope], ["准则", w.mandate.accounting_basis],
            *[["方法与适用范围", text] for text in data["mandate"]["methods"]],
            ["单位", "Facts D=原始值，E=倍数，F=基础单位。Calculations D=基础单位，E=展示值。"],
            ["复算", "数值公式可编辑；修改主体、期间、币种、规则或资料后须重跑 export.py 复核口径。"],
            ["勾稽精度", "金额残差按本次Decimal输入/计算保留的小数位ROUND，消除Excel二进制尾差，不改变业务容差；提高输入小数精度后须重跑导出。"],
            ["缺口", "空值不是零。未执行程序、假设与已披露数据分别标识。"],
            *[["限制", x] for x in data["mandate"]["limitations"]]], [22, 110])
        rows = []
        for f in w.facts:
            c = f.context
            rows.append([f.id, f.label, ", ".join(f.evidence), float(f.value) if f.value is not None else None,
                         float(c.scale), None, c.measure, c.currency, c.entity, c.scope,
                         str(c.start or ""), str(c.end), c.aggregation, f.state, c.basis, f.note, f.concept, c.physical_unit])
        ws = sheet("Facts", ["ID", "原始科目", "证据", "原始数值", "倍数", "基础单位数值", "量纲", "币种",
                             "实体", "范围", "开始", "结束", "统计类型", "状态", "准则", "说明", "概念", "实物单位"], rows)
        cells = {}
        records = {r.id: r for r in [*w.facts, *w.calculations]}
        for i, f in enumerate(w.facts, 2):
            cells[f.id] = f"'Facts'!F{i}"
            ws.write_formula(i-1, 5, f"=IF(COUNT(D{i},E{i})=2,D{i}*E{i},NA())", number,
                             float(f.value*f.context.scale) if f.value is not None else "#N/A")
        rs = {r["id"]: r for r in data["results"]["calculations"]}
        calcrows = [[c.id, c.label, "", None, None, rs[c.id]["status"], rs[c.id]["reason"],
                     c.definition, c.interpretation, unit(c.context), float(c.context.scale),
                     c.context.entity, c.context.scope, str(c.context.start or ""), str(c.context.end)] for c in w.calculations]
        ws = sheet("Calculations", ["ID", "指标或桥", "基础单位公式", "基础单位结果", "展示结果", "状态", "不计算原因",
                                    "定义", "解释边界", "展示单位", "倍数", "实体", "范围", "开始", "结束"], calcrows,
                   [14, 20, 44, 24, 24, 13, 18, 22, 24, 12, 8, 16, 10, 11, 11])
        for i, c in enumerate(w.calculations, 2):
            r = rs[c.id]
            expr = excel_expression(c, cells) if len(c.terms) == 2 or c.op == "sum" else "=NA()"
            if c.op in ("ratio", "growth") and len(c.terms) == 2:
                den = cells[c.terms[1].ref]
                test = f"{den}<=0" if c.denominator == "positive" or c.op == "growth" else f"{den}=0"
                if c.op == "growth":
                    test = f"OR({test},{cells[c.terms[0].ref]}<0)"
                expr = f"=IF({test},NA(),{expr[1:]})"
            ws.write_string(i-1, 2, expr, wrap)
            if not context_error(c, [records[t.ref] for t in c.terms]):
                ws.write_formula(i-1, 3, expr, number, float(r["normalized"]) if r["normalized"] is not None else "#N/A")
                if c.context.measure == "ratio":
                    is_points = c.context.physical_unit == "percentage_points"
                    ws.write_formula(i-1, 4, f"=D{i}*100" if is_points else f"=D{i}", points if is_points else times if c.context.physical_unit == "times" else percent,
                                     float(r["normalized"])*(100 if is_points else 1) if r["normalized"] is not None else "#N/A")
                else:
                    ws.write_formula(i-1, 4, f"=D{i}/K{i}", number, float(r["value"]) if r["value"] is not None else "#N/A")
                ws.write_formula(i-1, 5, f'=IF(ISNUMBER(D{i}),"calculated","not_calculated")', wrap, r["status"])
                ws.write_formula(i-1, 6, f'=IF(ISNUMBER(D{i}),"","检查缺失值、分母与输入；口径变更须重跑Python")', wrap,
                                 "" if r["normalized"] is not None else "检查缺失值、分母与输入；口径变更须重跑Python")
            else:
                ws.write_formula(i-1, 3, "=NA()", number, "#N/A")
                ws.write_formula(i-1, 4, "=NA()", number, "#N/A")
            cells[c.id] = f"'Calculations'!D{i}"
        checks = {r["id"]: r for r in data["results"]["reconciliations"]}
        normalized = {f.id: f.value*f.context.scale if f.value is not None else None for f in w.facts}
        normalized.update({key: Decimal(r["normalized"]) if r["normalized"] is not None else None for key, r in rs.items()})
        rows = [[r.id, r.label, r.actual, r.expected, None, float(r.tolerance), checks[r.id]["status"], r.basis, checks[r.id]["reason"]] for r in w.reconciliations]
        ws = sheet("Reconciliations", ["ID", "勾稽", "实际", "目标", "残差 基础单位", "容差 基础单位", "状态", "容差依据", "未测试原因", "金额残差小数位"], rows,
                   [18, 30, 22, 22, 24, 24, 25, 38, 28, 12])
        for i, r in enumerate(w.reconciliations, 2):
            value = checks[r.id]["residual"]
            if reconciliation_matches(records[r.actual], records[r.expected]):
                expr = f"{cells[r.actual]}-{cells[r.expected]}"
                amounts = [normalized[r.actual], normalized[r.expected]]
                if records[r.actual].context.measure == "money" and all(v is not None for v in amounts):
                    places = max(0, *(-v.as_tuple().exponent for v in amounts))
                    ws.write_number(i-1, 9, places)
                    expr = f"ROUND({expr},J{i})"
                ws.write_formula(i-1, 4, "="+expr, wrap, float(value) if value is not None else "#N/A")
                ws.write_formula(i-1, 6, f'=IF(COUNT(E{i},F{i})<2,"not_tested",IF(ABS(E{i})<=F{i},"within_input_tolerance","unexplained_difference"))', wrap, checks[r.id]["status"])
                ws.write_formula(i-1, 8, f'=IF(ISNUMBER(E{i}),"","missing or unavailable input")', wrap, checks[r.id]["reason"])
        sheet("Sources", ["ID", "文件", "链接", "公布日期", "SHA256", "时间证据与限制"],
              [[s.id, s.title, s.url, str(s.published) if s.published else "未核验", s.sha256, s.availability_note] for s in w.sources], [20, 60, 100, 20, 70, 70])
        sheet("Evidence", ["ID", "来源", "定位", "观察", "可靠性与限制"],
              [[e.id, e.source, e.locator, e.observation, e.reliability] for e in w.evidence], [20, 20, 45, 90, 65])
        sheet("Findings", ["ID", "问题", "结论", "机制", "状态", "依据", "反证", "其他解释", "改变结论条件"],
              [[f["id"], f["question"], f["conclusion"], f["mechanism"], f["status"], ", ".join(f["evidence"]),
                ", ".join(f["counterevidence"]), "\n".join(f["alternatives"]), f["changes_if"]] for f in data["findings"]], [18,45,85,75,18,35,35,75,75])
        sheet("Procedures", ["ID", "事项", "目的与认定", "总体", "选取", "步骤", "状态", "实际结果", "证据", "执行者日期"],
              [[p["id"], p["finding"], p["purpose"]+" / "+", ".join(p["assertions"]), p["population"], p["selection"],
                "\n".join(p["steps"]), p["status"], p["result"], ", ".join(p["evidence"]), f"{p['performed_by'] or ''} {p['performed_on'] or ''}"] for p in data["procedures"]])
        sheet("Requests", ["ID", "事项", "所需资料", "影响", "责任角色", "关闭条件"],
              [[r["id"], r["finding"], r["request"], r["reason"], r["owner_role"], r["close_when"]] for r in data["requests"]], [18,18,75,70,30,75])
        for index, q in enumerate(w.quantitative, 1):
            if q.method == "operating_cash_scenario_v2":
                scenario_sheets(book, sheet, q.artifact, f"Q{index}", number)
            elif q.method == "pit_margin_persistence_v1":
                panel_sheets(sheet, q.artifact, f"Q{index}", number, percent)
            else:
                rows = q.artifact.get("rows", [])
                keys = list(dict.fromkeys(k for row in rows for k in row))
                sheet(f"Quant{index}", keys or ["result"], [[row.get(k, "") for k in keys] for row in rows])
                if q.method == "single_entity_recovery_waterfall":
                    claims = {c["id"]: c for c in q.artifact["input_snapshot"]["claims"]}
                    totals = [r for r in rows if r["kind"] == "claim_total"]
                    ws = sheet(f"Q{index}Recovery", ["债权", "原金额", "担保池回收", "一般财产回收", "合计回收", "未偿", "回收比例"],
                               [[claims[r["claim_id"]]["label"], *[float(r[k]) for k in ("original_claim", "secured", "general", "result", "unpaid")],
                                 float(r["recovery_rate"]) if r["recovery_rate"] is not None else None] for r in totals])
                    for i, r in enumerate(totals, 2):
                        ws.write_formula(i-1, 4, f"=IF(COUNT(C{i}:D{i})=2,C{i}+D{i},NA())", number, float(r["result"]))
                        ws.write_formula(i-1, 5, f"=IF(COUNT(B{i},E{i})=2,B{i}-E{i},NA())", number, float(r["unpaid"]))
                        ws.write_formula(i-1, 6, f"=IF(AND(COUNT(B{i},E{i})=2,B{i}>0),E{i}/B{i},NA())", percent,
                                         float(r["recovery_rate"]) if r["recovery_rate"] is not None else "#N/A")
                    ws.merge_range(len(totals)+3, 0, len(totals)+4, 6, "回收分配是已运行结果快照；这里只联动合计、未偿和比例。修改估值、债权或顺位后，重跑 recovery.py 与 export.py，不在本表重新分配。", wrap)
            sheet(f"Quant{index}Notes", ["field", "value"], [["id", q.id], ["method", q.method], ["as_of", str(q.as_of)],
                  ["evidence", ", ".join(q.evidence)],
                  ["input references", ", ".join(q.input_refs)], ["assumptions", "\n".join(data["quantitative"][index-1]["assumptions"])],
                  ["limitations", "\n".join(data["quantitative"][index-1]["limitations"])], ["reproduction", f"See quantitative-{q.id}.json input_snapshot and method version"],
                  *([["execution", "已运行专门结果快照；本次导出未重新计算或校验。输入或方法改变后须重跑专门脚本并重新导出。"]]
                    if q.method not in ("operating_cash_scenario_v2", "pit_margin_persistence_v1", "single_entity_recovery_waterfall") else []),
                  *[[f.id, f"{f.label}；artifact.rows[{f.row}].{f.field}；{f.context.model_dump_json()}"] for f in q.figures]], [28,110])


def scenario_sheets(book, sheet, artifact, prefix, number):
    snapshot = artifact["input_snapshot"]
    opening, floor = snapshot["opening"], snapshot["minimum_cash"]
    sheet(prefix+"Inputs", ["输入", "值", "依据"], [
        ["amount_scale", snapshot["amount_scale"], snapshot["currency"]],
        ["interest_basis_days", snapshot["interest_basis_days"], "实际期间天数 / 此日数"],
        ["payables_denominator", snapshot["payables_denominator"], "purchases或cost_of_sales"],
        ["minimum_cash", floor["value"], floor["source"]],
        *[[key, opening[key], opening["source"]] for key in ("cash", "receivables", "inventory", "payables", "debt")],
        ["复算边界", "编辑合法数字驱动可重算；缺失、非数字或越界驱动为#N/A。日期/来源/口径/期间或逆向目标改变须重跑脚本", "负现金后续期间为#N/A；不自动融资"]], [32,55,85])
    inputs = f"'{prefix}Inputs'!"
    driver_keys = ["start", "end", "days", "volume", "unit_price", "unit_cost_of_sales", "fixed_cash_cost",
                   "depreciation", "capex", "tax_rate", "dso", "dio", "dpo", "interest_rate", "drawdown", "principal", "dividends",
                   "depreciation_in_cost_of_sales", "inventory_cash_conversion", "inventory_depreciation_change", "source", "available_at"]
    drivers = sheet(prefix+"Drivers", driver_keys, [[p.get(k, "") for k in driver_keys] for p in snapshot["periods"]])
    datefmt = book.add_format({"num_format": "yyyy-mm-dd"})
    for i, period in enumerate(snapshot["periods"], 2):
        start, end = datetime.fromisoformat(period["start"]), datetime.fromisoformat(period["end"])
        drivers.write_datetime(i-1, 0, start, datefmt)
        drivers.write_datetime(i-1, 1, end, datefmt)
        drivers.write_formula(i-1, 2, f"=B{i}-A{i}+1", number, (end-start).days+1)
    keys = ["period_end", "revenue", "cost_of_sales", "ebitda", "depreciation", "ebit", "cash_interest", "cash_taxes", "net_income",
            "receivables_end", "inventory_end", "implied_purchases", "payables_end", "nwc_end", "delta_nwc", "cfo",
            "capex", "drawdown", "principal", "dividends", "cash_begin", "cash_end", "debt_begin", "debt_end",
            "cash_headroom", "funding_needed_to_floor", "interest_coverage", "debt_to_ebitda", "status",
            "customer_collections", "supplier_payments", "inventory_depreciation_lower", "inventory_depreciation_upper"]
    cash = sheet(prefix+"Cash", keys, [])
    cash.autofilter(0, 0, len(snapshot["periods"]), len(keys)-1)
    cash.set_header(f"&L{prefix}Cash&R{snapshot['currency']} × {snapshot['amount_scale']:g}")
    output = {r["period_end"]: dict(r) for r in artifact["rows"]}
    ar, ap, low, high = opening["receivables"], opening["payables"], 0, opening["inventory"]
    for period in snapshot["periods"]:
        row = output.get(period["end"])
        if row is None:
            break
        low = max(0, low + period["inventory_depreciation_change"])
        high = min(row["inventory_end"], high + period["inventory_depreciation_change"])
        row.update(customer_collections=max(0, ar + row["revenue"] - row["receivables_end"]),
                   supplier_payments=max(0, ap + row["implied_purchases"] - row["payables_end"]),
                   inventory_depreciation_lower=low, inventory_depreciation_upper=high)
        ar, ap = row["receivables_end"], row["payables_end"]
    for i, period in enumerate(snapshot["periods"], 2):
        d = lambda col: f"'{prefix}Drivers'!{col}{i}"
        valid = (f"IFERROR(AND(COUNT({inputs}B2:B3,{inputs}B5:B10)=8,{inputs}B2>0,{inputs}B3>0,"
                 f"MIN({inputs}B5:B10)>=0,OR({inputs}B4=\"purchases\",{inputs}B4=\"cost_of_sales\"),"
                 f"COUNT('{prefix}Drivers'!A{i}:T{i})=20,{d('C')}>0,"
                 f"MIN('{prefix}Drivers'!D{i}:M{i},'{prefix}Drivers'!O{i}:S{i})>=0,{d('J')}<=1,"
                 f"{d('R')}<=MIN({d('H')},{d('D')}*{d('F')}),{d('R')}+{d('T')}>=0),FALSE)")
        prev = i-1
        previous_cash = f"V{prev}" if i > 2 else inputs+"B6"
        previous_debt = f"X{prev}" if i > 2 else inputs+"B10"
        previous_inventory = f"K{prev}" if i > 2 else inputs+"B8"
        previous_nwc = f"N{prev}" if i > 2 else f"({inputs}B7+{inputs}B8-{inputs}B9)"
        previous_ar = f"J{prev}" if i > 2 else inputs+"B7"
        previous_ap = f"M{prev}" if i > 2 else inputs+"B9"
        previous_low = f"AF{prev}" if i > 2 else "0"
        previous_high = f"AG{prev}" if i > 2 else inputs+"B8"
        collectible, payable = f"({previous_ar}+B{i})", f"({previous_ap}+L{i})"
        feasible = f"AND(COUNT(AD{i}:AG{i})=4,AF{i}<=AG{i})"
        formulas = [f"{d('D')}*{d('E')}", f"{d('D')}*{d('F')}", f"B{i}-C{i}+{d('R')}-{d('G')}", d('H'),
                    f"D{i}-E{i}", f"W{i}*{d('N')}*{d('C')}/{inputs}B3", f"MAX(F{i}-G{i},0)*{d('J')}",
                    f"F{i}-G{i}-H{i}", f"B{i}*{d('K')}/{d('C')}", f"C{i}*{d('L')}/{d('C')}",
                    f"C{i}+K{i}-{previous_inventory}-{d('S')}-{d('R')}-{d('T')}",
                    f'IF(L{i}<0,NA(),IF({inputs}B4="purchases",L{i},C{i}))*{d("M")}/{d("C")}',
                    f"IF({feasible},J{i}+K{i}-M{i},NA())", f"N{i}-{previous_nwc}", f"I{i}+E{i}+{d('T')}-O{i}", d('I'), d('O'), d('P'), d('Q'),
                    previous_cash, f"IF(ISNUMBER(X{i}),U{i}+P{i}-Q{i}+R{i}-S{i}-T{i},NA())", previous_debt,
                    f"IF(W{i}+R{i}-S{i}<0,NA(),W{i}+R{i}-S{i})", f"V{i}-{inputs}B5", f"MAX({inputs}B5-V{i},0)",
                    f"IF(G{i}>0,F{i}/G{i},NA())", f"IF(D{i}>0,X{i}/D{i},NA())"]
        flow_formulas = [
            f"IF(J{i}-{collectible}>1E-14*MAX(ABS(J{i}),ABS({collectible})),NA(),MAX(0,{collectible}-J{i}))",
            f"IF(M{i}-{payable}>1E-14*MAX(ABS(M{i}),ABS({payable})),NA(),MAX(0,{payable}-M{i}))",
            f"MAX(0,{previous_low}+{d('T')})", f"MIN(K{i},{previous_high}+{d('T')})"]
        row = output.get(period["end"], {})
        cash.write(i-1, 0, period["end"])
        for col, formula in [*enumerate(formulas, 1), *enumerate(flow_formulas, 29)]:
            if i > 2:
                formula = f"IF({previous_cash}<0,NA(),{formula})"
            formula = f"IF({valid},{formula},NA())"
            value = row.get(keys[col])
            cash.write_formula(i-1, col, "="+formula, number, value if value is not None else "#N/A")
        cash.write_formula(i-1, 28, f'=IF(NOT({valid}),"invalid_numeric_inputs",IF(IFERROR({previous_cash}>=0,FALSE),IF(IFERROR({feasible},FALSE),IFERROR(IF(V{i}<0,"unfunded_cash_shortfall",IF(V{i}<{inputs}B5,"below_cash_floor","conditional")),"not_projected"),"infeasible_cash_path"),"not_projected"))',
                           None, row.get("status", "not_projected"))
    print_start = len(snapshot["periods"])+4
    print_row = print_start
    label_format = book.add_format({"bold": True, "text_wrap": True, "valign": "top"})
    print_number = book.add_format({"num_format": "#,##0.00;[Red](#,##0.00)", "valign": "top"})
    for offset in range(0, len(snapshot["periods"]), 4):
        periods = snapshot["periods"][offset:offset+4]
        cash.write(print_row, 0, "现金路径打印视图 / Cash path", label_format)
        cash.set_row(print_row, 30)
        cash.write_row(print_row, 1, [p["end"] for p in periods], label_format)
        for k, key in enumerate(keys[1:], print_row+1):
            cash.write(k, 0, key, label_format)
            cash.set_row(k, 20)
            for column, period in enumerate(periods, 1):
                value = output.get(period["end"], {}).get(key, "not_projected" if key == "status" else "#N/A")
                cash.write_formula(k, column, f"={xl_col_to_name(keys.index(key))}{offset+column+1}",
                                   label_format if key == "status" else print_number, value if value is not None else "#N/A")
        print_row += len(keys)+2
    cash.set_paper(9 if len(snapshot["periods"]) <= 2 else 8)
    if len(snapshot["periods"]) > 2:
        cash.set_landscape()
    cash.set_margins(left=0.25, right=0.25, top=0.6, bottom=0.35)
    cash.fit_to_pages(1, 0)
    cash.print_area(print_start, 0, print_row-3, min(4, len(snapshot["periods"])))
    cash.set_h_pagebreaks(list(range(print_start+len(keys)+2, print_row, len(keys)+2)))
    contracts = artifact.get("contracts", [])
    ws = sheet(prefix+"Contracts", ["条件", "日期", "指标", "关系", "阈值", "计算值", "结果", "定义", "依据"],
               [[c["label"], c["test_date"], c["metric"], c["relation"], c["threshold"], c["value"], c["status"], c["definition"], c["source"]] for c in contracts])
    dates = [p["end"] for p in snapshot["periods"]]
    for i, c in enumerate(contracts, 2):
        if c["test_date"] in dates:
            row = dates.index(c['test_date'])+2
            ref = f"'{prefix}Cash'!{xl_col_to_name(keys.index(c['metric']))}{row}"
            ws.write_formula(i-1, 5, f"=IF(ISNUMBER('{prefix}Cash'!V{row}),{ref},NA())", number, c["value"] if c["value"] is not None else "#N/A")
            comparison = ">=" if c["relation"] == "at_least" else "<="
            ws.write_formula(i-1, 6, f'=IF(COUNT(E{i}:F{i})=2,IF(F{i}{comparison}E{i},"within_input_threshold","outside_input_threshold"),"not_tested")', None, c["status"])
    reverse = artifact.get("reverse")
    if reverse:
        rows = [["status", reverse["status"]], ["definition", reverse["definition"]],
                ["recalculation", "逆向根由SciPy计算；改边界/目标后重跑，不自动再次求根"]]
        for key in ("driver", "period_index", "driver_value", "test_date", "target_cash", "cash_residual", "source", "reason"):
            if key in reverse:
                rows.append([key, reverse[key]])
        for label, value in zip(("lower_bound", "upper_bound"), reverse["bounds"]):
            rows.append([label, value])
        for label, value in zip(("lower_cash_residual", "upper_cash_residual"), reverse["endpoint_cash_residuals"]):
            rows.append([label, value])
        sheet(prefix+"Reverse", ["字段", "结果"], rows, [35,100])
        if reverse.get("rows"):
            summary = ["period_end", "revenue", "net_income", "cfo", "cash_end", "debt_end", "cash_headroom"]
            sheet(prefix+"AtBoundary", summary, [[r[k] for k in summary] for r in reverse["rows"]])


def panel_sheets(sheet, artifact, prefix, number, percent):
    records = artifact["input_snapshot"]["records"]
    keys = ["id", "entity", "period_end", "available_at", "version", "revenue", "cost_of_sales", "total_assets", "industry", "accounting_basis", "currency", "source"]
    sheet(prefix+"Records", keys, [[r[k] for k in keys] for r in records])
    locations = {r["id"]: i for i, r in enumerate(records, 2)}
    margin = lambda identifier: f"IF(AND(COUNT('{prefix}Records'!F{locations[identifier]}:G{locations[identifier]})=2,'{prefix}Records'!F{locations[identifier]}>0),('{prefix}Records'!F{locations[identifier]}-'{prefix}Records'!G{locations[identifier]})/'{prefix}Records'!F{locations[identifier]},NA())"
    peers = [r for r in artifact["rows"] if r["kind"] == "peer"]
    ws = sheet(prefix+"Peers", ["企业", "选用记录", "可用时间", "版本", "毛利率", "毛利/期末资产", "样本内百分位", "有效同业数"],
               [[r["entity"], r["record_id"], r["available_at"], r["version"], r["gross_margin"], r["gross_profit_to_assets"], r["gross_profitability_percentile"], r["peer_n"]] for r in peers])
    for i, row in enumerate(peers, 2):
        location = locations[row["record_id"]]
        ws.write_formula(i-1, 4, "=IFERROR("+margin(row["record_id"])+",NA())", percent, row["gross_margin"] if row["gross_margin"] is not None else "#N/A")
        expr = f"('{prefix}Records'!F{location}-'{prefix}Records'!G{location})/'{prefix}Records'!H{location}"
        ws.write_formula(i-1, 5, f"=IF(AND(COUNT('{prefix}Records'!F{location}:H{location})=3,'{prefix}Records'!H{location}>0),{expr},NA())", percent, row["gross_profit_to_assets"] if row["gross_profit_to_assets"] is not None else "#N/A")
        area = f"F$2:F${len(peers)+1}"
        ws.write_formula(i-1, 6, f'=IF(ISNUMBER(F{i}),IFERROR((COUNTIF({area},"<"&F{i})+(COUNTIF({area},F{i})+1)/2)/COUNT({area}),NA()),NA())', percent,
                         row["gross_profitability_percentile"] if row["gross_profitability_percentile"] is not None else "#N/A")
        ws.write_formula(i-1, 7, f"=COUNT({area})", number, row["peer_n"])
    train = artifact["training_pairs"]
    ws = sheet(prefix+"Train", ["企业", "首次起点记录", "首次目标记录", "当前毛利率", "下一年毛利率"],
               [[p["entity"], p["origin_id"], p["target_id"], p["x"], p["actual"]] for p in train])
    for i, pair in enumerate(train, 2):
        ws.write_formula(i-1, 3, "="+margin(pair["origin_id"]), percent, pair["x"])
        ws.write_formula(i-1, 4, "="+margin(pair["target_id"]), percent, pair["actual"])
    model = artifact["model"]
    ws = sheet(prefix+"Model", ["字段", "结果"], [["training_n", model["training_n"]], ["test_n", model["test_n"]],
               ["intercept", model.get("intercept")], ["slope", model.get("slope")], ["test_mae", model.get("test_mae")],
               ["baseline_mae", model.get("baseline_mae")], ["status", model["status"]],
               ["解释", "样本/时间/版本改变须重跑panel.py；固定样本数值变化可复算OLS及留出误差。不是PD或显著性证明。"]], [32,110])
    if model["status"] == "estimated":
        last = len(train)+1
        for excel_row, function, value in [(4, "INTERCEPT", model["intercept"]), (5, "SLOPE", model["slope"])]:
            ws.write_formula(excel_row-1, 1, f"=IFERROR({function}('{prefix}Train'!E2:E{last},'{prefix}Train'!D2:D{last}),NA())", number, value)
    predictions = [r for r in artifact["rows"] if r["kind"] == "out_of_time_prediction"]
    predsheet = sheet(prefix+"Test", ["企业", "起点记录", "目标记录", "当前值", "实际下一期", "预测", "预测绝对误差", "Last-value绝对误差", "起点可用时间", "目标可用时间"],
                      [[p["entity"], p["origin_id"], p["target_id"], p["x"], p["actual"], p["predicted"], abs(p["predicted"]-p["actual"]), abs(p["x"]-p["actual"]), p["origin_available_at"], p["target_available_at"]] for p in predictions])
    for i, p in enumerate(predictions, 2):
        for col, identifier, value in [(3, p["origin_id"], p["x"]), (4, p["target_id"], p["actual"])]:
            predsheet.write_formula(i-1, col, "="+margin(identifier), percent, value)
        predsheet.write_formula(i-1, 5, f"='{prefix}Model'!B4+'{prefix}Model'!B5*D{i}", percent, p["predicted"])
        predsheet.write_formula(i-1, 6, f"=ABS(F{i}-E{i})", percent, abs(p["predicted"]-p["actual"]))
        predsheet.write_formula(i-1, 7, f"=ABS(D{i}-E{i})", percent, abs(p["x"]-p["actual"]))
    if predictions:
        for excel_row, col, value in [(6, "G", model["test_mae"]), (7, "H", model["baseline_mae"])]:
            ws.write_formula(excel_row-1, 1, f"=AVERAGE('{prefix}Test'!{col}2:{col}{len(predictions)+1})", percent, value)
    sheet(prefix+"Selection", ["类别", "记录", "原因"],
          [[kind, p["record_id"], p["reason"]] for kind, pairs in [("pending", artifact["pending_outcomes"]), ("excluded", artifact["excluded_pairs"])] for p in pairs], [22,45,100])


def word(data, path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    footer = section.footer.paragraphs[0]
    footer.text = f"{data['mandate']['entity']}  |  {data['mandate']['version']}  |  "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.autofit = False
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:"+edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                border.set(qn("w:"+key), value)
            borders.append(border)
        t._tbl.tblPr.append(borders)
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = text
            cell.paragraphs[0].paragraph_format.keep_with_next = True
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E6EEF0")
            cell._tc.get_or_add_tcPr().append(shading)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, text in zip(t.add_row().cells, row):
                cell.text = str(text)
        weights = [.35, .45, .20] if len(headers) == 3 else [1/len(headers)]*len(headers)
        available_width = section.page_width - section.left_margin - section.right_margin
        for col, weight in zip(t.columns, weights):
            col.width = int(available_width*weight)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, weight in zip(row.cells, weights):
                cell.width = int(available_width*weight)
        doc.add_paragraph()

    m = data["mandate"]
    doc.add_paragraph(m["title"], "Title")
    doc.add_paragraph(f"{m['period_start']} — {m['period_end']}　资料截止 {m['cutoff']}　底稿 {m['version']}")
    doc.add_paragraph(m["purpose"])
    for text in m["methods"]:
        doc.add_paragraph("方法与适用范围："+text)
    findings = {f["id"]: f for f in data["findings"]}
    quantitative_origins = {f["id"]: q["id"] for q in data["quantitative"] for f in q["figures"]}
    if data["findings"]:
        doc.add_heading("核心判断", 1)
    for f in data["findings"][:3]:
        doc.add_paragraph(f"[{f['status']}] {f['conclusion']} [{f['id']}; {', '.join(f['evidence'])}]")
    doc.add_paragraph(f"范围：{m['scope']}。会计基础：{m['accounting_basis']}。")
    for text in m["limitations"]:
        doc.add_paragraph(text)
    for s in data["sections"]:
        doc.add_heading(s["title"], 1)
        if s["figures"]:
            table(["项目及记录", "期间及数值", "证据"],
                  [[data['figures'][ref]['label']+f" [{ref}]", data['figures'][ref]['context']['end']+"\n"+data['figures'][ref]['display'],
                    "专门结果 "+quantitative_origins[ref]+"（输入与证据见该节）" if ref in quantitative_origins
                    else ", ".join(data['figures'][ref]['evidence'])] for ref in s["figures"]])
        for ref in s["findings"]:
            f = findings[ref]
            doc.add_heading(f["title"]+f" [{ref}; {f['status']}]", 2)
            for text in (f["conclusion"], f["mechanism"], "证据："+", ".join(f["evidence"]),
                         "反证："+(", ".join(f["counterevidence"]) or "尚无足以排除其他解释的额外证据"),
                         "其他解释："+"；".join(f["alternatives"]), "判断改变条件："+f["changes_if"]):
                doc.add_paragraph(text)
    if data["reconciliations"]:
        doc.add_heading("勾稽结果", 1)
    for r in data["reconciliations"]:
        result = next(x for x in data["results"]["reconciliations"] if x["id"] == r["id"])
        doc.add_paragraph(f"{r['label']} [{r['id']}]：{result['status']}；基础单位残差 {result['residual'] if result['residual'] is not None else '未计算'}。{result['reason']} {r['basis']}")
    for index, q in enumerate(data["quantitative"], 1):
        doc.add_page_break()
        doc.add_heading(q["label"]+f" [{q['id']}]", 1)
        doc.add_paragraph(f"{q['method']}；截至 {q['as_of']}。")
        artifact = q["artifact"]
        rows = artifact.get("rows", [])
        fmt = lambda value: "未计算" if value is None else f"{value:,.4f}" if isinstance(value, (int, float)) else str(value)
        if q["method"] == "operating_cash_scenario_v2":
            minimum = min(rows, key=lambda row: row["cash_end"])
            doc.add_paragraph(f"金额单位 {artifact['currency']} × {artifact['amount_scale']:g}。最低期末现金 {minimum['cash_end']:,.4f}，发生在 {minimum['period_end']}；补至输入最低现金所需资金 {minimum['funding_needed_to_floor']:,.4f}。未自动补融资。")
            table(["期间", "收入", "净利润", "CFO", "期末现金", "期末债务"],
                  [[r["period_end"], *[fmt(r[k]) for k in ("revenue", "net_income", "cfo", "cash_end", "debt_end")]] for r in rows])
            table(["期间", "营运资本增加", "资本开支", "借入", "还本", "股利"],
                  [[r["period_end"], *[fmt(r[k]) for k in ("delta_nwc", "capex", "drawdown", "principal", "dividends")]] for r in rows])
            if artifact.get("contracts"):
                doc.add_heading("输入合同条件", 2)
                table(["条件和日期", "定义与阈值", "计算结果"],
                      [[c["label"]+"\n"+c["test_date"], f"{c['definition']}\n{c['metric']} {c['relation']} {fmt(c['threshold'])}\n来源 {c['source']}",
                        f"{fmt(c['value'])}\n{c['status']}"] for c in artifact["contracts"]])
            if artifact.get("reverse"):
                reverse = artifact["reverse"]
                doc.add_heading("逆向现金边界", 2)
                doc.add_paragraph(reverse["definition"]+"；"+reverse["status"])
                if reverse["status"] == "converged":
                    table(["变量与期间", "边界与目标", "求解残差"],
                          [[f"{reverse['driver']}\n{reverse['test_date']}", f"驱动值 {reverse['driver_value']:,.6f}\n目标现金 {reverse['target_cash']:,.4f}", f"{reverse['cash_residual']:.8g}"]])
                    doc.add_paragraph("这是输入假设下的现金边界，不是发生概率或法律违约判断。")
                else:
                    doc.add_paragraph(reverse.get("reason", "未取得收敛边界"))
            if artifact["unprojected_periods"]:
                doc.add_paragraph(f"因现金未融资缺口，后续 {artifact['unprojected_periods']} 期未继续预测。")
        elif q["method"] == "pit_margin_persistence_v1":
            model = artifact["model"]
            table(["检验", "结果", "解释"], [["训练与留出", f"{model['training_n']} / {model['test_n']}", model['status']],
                  ["留出MAE", fmt(model.get("test_mae")), "毛利率小数单位；不是训练准确率"],
                  ["Last-value基准MAE", fmt(model.get("baseline_mae")), "使用上一期毛利率预测下一期"]])
            peers = [r for r in rows if r["kind"] == "peer"]
            table(["企业", "选用版本", "毛利率", "毛利资产比", "样本内百分位"],
                  [[r["entity"], r["version"], fmt(r["gross_margin"]), fmt(r["gross_profit_to_assets"]), fmt(r["gross_profitability_percentile"])] for r in peers])
            predictions = [r for r in rows if r["kind"] == "out_of_time_prediction"]
            if predictions:
                table(["企业", "预测", "实际下一期", "Last-value基准"],
                      [[p["entity"], fmt(p["predicted"]), fmt(p["actual"]), fmt(p["last_value_baseline"])] for p in predictions])
            doc.add_paragraph(f"未成熟结果 {len(artifact['pending_outcomes'])} 条；其他排除 {len(artifact['excluded_pairs'])} 条。记录、版本和排除理由详见工作簿。")
        elif q["method"] == "single_entity_recovery_waterfall":
            snapshot = artifact["input_snapshot"]
            claims = {c["id"]: c for c in snapshot["claims"]}
            doc.add_paragraph(f"单一法人 {snapshot['entity']}；金额单位 {snapshot['currency']} {snapshot['unit']}。估值基础：{snapshot['value_basis']}。顺位依据：{snapshot['priority_basis']}。")
            table(["债权", "原金额", "担保池回收", "一般回收", "总回收", "未偿", "回收比例"],
                  [[claims[r["claim_id"]]["label"], r["original_claim"], r["secured"], r["general"], r["result"], r["unpaid"],
                    f"{Decimal(r['recovery_rate']):.2%}" if r["recovery_rate"] is not None else "未定义"]
                   for r in rows if r["kind"] == "claim_total"])
            for r in rows:
                if r["kind"] in ("cost", "estate_residual", "reconciliation"):
                    doc.add_paragraph(f"{r['kind']} / {r['pool']}：{r['result']}；{r['formula']}；输入 {r['raw_input']}。")
            doc.add_paragraph("分配顺位、估值、抵押池或债权改变必须重跑 recovery.py。工作簿只对已分配回收的合计、未偿与比例提供联动公式，不重新决定法律顺位。")
        else:
            doc.add_paragraph("已运行专门结果快照；本次导出未重新计算或校验。输入或方法改变后须重跑专门脚本并重新导出。")
            if q["figures"]:
                table(["指标", "结果及单位", "底稿引用"],
                      [[f["label"], data["figures"][f["id"]]["display"], f["id"]] for f in q["figures"]])
                doc.add_paragraph("此处仅汇总声明的指标；全部结果行、公式和未舍入值保留在工作簿及结果JSON。")
            else:
                for row_index, row in enumerate(rows):
                    doc.add_heading(f"结果行 {row_index}", 2)
                    table(["字段", "值"], [[key, "null" if value is None else
                          json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)]
                          for key, value in row.items()])
            if not rows:
                doc.add_paragraph("未提供 artifact.rows 结果行。")
        doc.add_paragraph("证据记录："+", ".join(q["evidence"]))
        doc.add_heading("假设与适用限制", 2)
        for text in q["assumptions"]+q["limitations"]:
            doc.add_paragraph(text)
        doc.add_paragraph(f"完整结果、输入快照与复算信息见 Excel Q{index}/Quant{index} 系列工作表及 quantitative-{q['id']}.json。")
    if data["procedures"] or data["requests"]:
        doc.add_heading("已执行程序与待核工作", 1)
    for p in data["procedures"]:
        doc.add_heading(f"{p['id']} {p['purpose']}", 2)
        doc.add_paragraph(f"状态 {p['status']}；认定 {'、'.join(p['assertions'])}；总体 {p['population']}；选取 {p['selection']}。")
        for step in p["steps"]:
            doc.add_paragraph(step)
        doc.add_paragraph(f"实际结果：{p['result']}；证据 {', '.join(p['evidence'])}。")
    for r in data["requests"]:
        doc.add_heading(f"{r['id']} 待补资料", 2)
        doc.add_paragraph(r["request"])
        doc.add_paragraph(f"影响：{r['reason']}。责任角色：{r['owner_role']}。关闭条件：{r['close_when']}")
    doc.add_heading("来源与定位", 1)
    for s in data["sources"]:
        doc.add_paragraph(f"[{s['id']}] {s['title']}；公布 {s['published'] or '未核验'}\n{s['url']}\n{s['availability_note']}")
    for e in data["evidence"]:
        doc.add_paragraph(f"[{e['id']}] {e['source']} {e['locator']}。{e['observation']} {e['reliability']}")
    for border in doc.element.xpath(".//w:pPr/w:pBdr"):
        border.getparent().remove(border)
    for text_node in doc.element.iter(qn("w:t")):
        text_node.text = re.sub(r"(?<![0-9A-Za-z_./])[−-](?=\d)", "\u2011", text_node.text)
    doc.save(path)


def export(input_path, output_dir, node="node"):
    raw = Path(input_path).read_bytes()
    w = Workpaper.model_validate_json(raw)
    result = evaluate(w)
    data = prepare(w, result)
    out = Path(output_dir).resolve()
    skill_dir = Path(__file__).resolve().parents[1]
    if out == skill_dir or skill_dir in out.parents:
        raise ValueError("output must be in the task workspace, outside the installed skill")
    out.mkdir(parents=True, exist_ok=True)
    (out / "workpaper.json").write_text(w.model_dump_json(indent=2), encoding="utf-8")
    (out / "results.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "presentation-data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    for q in w.quantitative:
        (out / f"quantitative-{q.id}.json").write_text(json.dumps(q.artifact, ensure_ascii=False, indent=2), encoding="utf-8")
    workbook(w, data, out / "workbook.xlsx")
    word(data, out / "report.docx")
    presentation = subprocess.run([node, str(Path(__file__).with_name("export_pptx.cjs")), str(out / "presentation-data.json"), str(out / "presentation.pptx")], check=True, stdout=subprocess.PIPE, text=True)
    manifest = {"version": w.mandate.version, "input_sha256": hashlib.sha256(raw).hexdigest(),
                "export_environment": {"python": platform.python_version(),
                    **{name: version(name) for name in ("pydantic", "python-docx", "XlsxWriter")},
                    **json.loads(presentation.stdout)},
                "scripts_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in Path(__file__).parent.iterdir() if p.suffix in (".py", ".cjs")},
                "files": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in out.iterdir() if p.name in
                          ("report.docx", "workbook.xlsx", "presentation.pptx", "workpaper.json", "results.json")}}
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return {"output": str(out), "calculations": len(result["calculations"]),
            "not_calculated": [r["id"] for r in result["calculations"] if r["value"] is None]}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input")
    parser.add_argument("--output", required=True)
    parser.add_argument("--node", default="node")
    args = parser.parse_args()
    print(json.dumps(export(args.input, args.output, args.node), ensure_ascii=False))
